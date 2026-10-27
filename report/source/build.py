"""Capstone report builder: markdown-like sources -> .docx (python-docx) -> PDF (LibreOffice), two passes so the
static table of contents / lists of figures and tables carry real page numbers."""
import os, re, sys, json, subprocess, shutil
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).parent
EDITION = os.environ.get("EDITION", "")
SIMPLE = EDITION == "simple"                       # plain-language edition with the team details
SHORT = EDITION in ("short", "simple")             # the condensed editions
REPO = HERE.parent.parent if (HERE.parent.parent / "src").exists() else Path("/home/user/RAG-Project")
SRC = HERE / ("src_simple" if SIMPLE else "src_short" if SHORT else "src")
OUT = HERE.parent if HERE.name == "source" else HERE / "out"
FONT = "Times New Roman"
MONO = "Courier New"
TEXT_W = 16.3 if SHORT else 15.5  # cm: A4 21 minus margins
BODY = 11 if SHORT else 12
LINE = 1.15 if SHORT else 1.5
SHORT_TITLE = "HF Model Finder: Constraint-Verified RAG for Hugging Face Models"
LOGO = HERE / "assets" / "logo.jpeg"

from refs import R as REFS

# ----------------------------------------------------------------------------- low-level helpers

def set_font(run, size=None, bold=None, italic=None, mono=False, color=None):
    name = MONO if mono else FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)

def para_fmt(p, align=None, line=1.5, before=0, after=6, left=None, first=None, keep_next=False, keep_lines=False, hang=None, widow=True):
    pf = p.paragraph_format
    if align is not None: p.alignment = align
    pf.line_spacing = line
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if left is not None: pf.left_indent = Cm(left)
    if first is not None: pf.first_line_indent = Cm(first)
    if hang is not None:
        pf.left_indent = Cm(hang[0]); pf.first_line_indent = Cm(-hang[1])
    pf.keep_with_next = keep_next
    pf.keep_together = keep_lines
    pf.widow_control = widow

def shade(el_pr, fill):
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    el_pr.append(shd)

def shade_par(p, fill):
    shade(p._p.get_or_add_pPr(), fill)

def shade_cell(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)

def cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)

def table_borders(table, color="8c8c8c", sz=4, inside=True):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    names = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for n in names:
        e = OxmlElement(f"w:{n}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color); b.append(e)
    tblPr.append(b)

def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for n in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{n}"); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)

def row_props(row, header=False, cant_split=True):
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))

def set_col_widths(table, widths_cm):
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for row in table.rows:
        for i, c in enumerate(row.cells):
            c.width = Cm(widths_cm[i])

def add_field(par, instr, size=None, bold=None):
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    rpr.append(rf)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = "1"; r.append(t)
    fld.append(r); par._p.append(fld)

def bottom_border(par, color="808080", sz=6, top=False):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:top" if top else "w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "4"); e.set(qn("w:color"), color)
    b.append(e); pPr.append(b)


# ----------------------------------------------------------------------------- schema ordering (Word is strict)
_ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
}
def fix_order(root):
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for tag, names in _ORDER.items():
        idx = {W + n: i for i, n in enumerate(names.split())}
        for el in root.iter(W + tag):
            kids = list(el)
            kids_sorted = sorted(kids, key=lambda k: idx.get(k.tag, 999))
            if kids != kids_sorted:
                for k in kids: el.remove(k)
                for k in kids_sorted: el.append(k)

# ----------------------------------------------------------------------------- inline markup

INLINE = re.compile(r"(\{\{.+?\}\}|\*\*.+?\*\*|(?<![\w*])\*(?!\s)[^*]+?(?<!\s)\*(?![\w*])|`[^`]+`|\[@[^\]]+\]|@(?:fig|tab|lst|alg|eq):[A-Za-z0-9_]+|<br>|\^\{[^}]+\}|_\{[^}]+\})")

class Ctx:
    def __init__(self):
        self.cites = []          # keys in order of first citation
        self.labels = {}         # 'fig:key' -> 'Fig. 3.2'
        self.headings = []       # (level, label, title)
        self.figs = []           # (label, caption)
        self.tabs = []
        self.collect_only = True

    def cite_no(self, key):
        if key not in REFS:
            raise KeyError(f"unknown reference key: {key}")
        if key not in self.cites:
            self.cites.append(key)
        return self.cites.index(key) + 1

def fmt_cites(keys, ctx):
    nums = sorted({ctx.cite_no(k) for k in keys})
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        if j - i >= 2: out.append(f"[{nums[i]}]–[{nums[j]}]")
        else: out.extend(f"[{n}]" for n in nums[i:j + 1])
        i = j + 1
    return ", ".join(out)

def add_inline(par, text, ctx, size=12, base_bold=False, base_italic=False, mono=False, color=None):
    pos = 0
    def emit(t, bold=base_bold, italic=base_italic, code=False, hl=False, sup=False, sub=False):
        if not t: return
        r = par.add_run(t)
        set_font(r, size - (1 if code else 0) if code else size, bold, italic, mono=(code or mono), color=color)
        if hl: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if sup: r.font.superscript = True
        if sub: r.font.subscript = True
    for m in INLINE.finditer(text):
        emit(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("{{"):
            emit(tok[2:-2], hl=True)
        elif tok.startswith("**"):
            add_inline(par, tok[2:-2], ctx, size, True, base_italic, mono, color)
        elif tok.startswith("*"):
            emit(tok[1:-1], italic=not base_italic)
        elif tok.startswith("`"):
            emit(tok[1:-1], code=True)
        elif tok.startswith("[@"):
            keys = [k.strip().lstrip("@") for k in tok[1:-1].split(";")]
            emit(fmt_cites(keys, ctx))
        elif tok.startswith("@"):
            emit(ctx.labels.get(tok[1:], "??" + tok[1:]), bold=False)
        elif tok == "<br>":
            par.add_run().add_break()
        elif tok.startswith("^{"):
            emit(tok[2:-1], sup=True)
        elif tok.startswith("_{"):
            emit(tok[2:-1], sub=True)
        pos = m.end()
    emit(text[pos:])

# ----------------------------------------------------------------------------- source parsing

def read_blocks(path):
    """Split a source file into blocks: heading, para, list, fence (::: ... :::), pagebreak."""
    lines = Path(path).read_text(encoding="utf-8").split("\n")
    blocks, i, buf = [], 0, []
    def flush():
        nonlocal buf
        if buf:
            blocks.append({"t": "para", "text": " ".join(s.strip() for s in buf)}); buf = []
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("<!--"):
            while "-->" not in lines[i]: i += 1
            i += 1; continue
        if ln.startswith("::: "):
            flush()
            head = ln[4:].strip(); body = []; i += 1
            while not lines[i].startswith(":::"):
                body.append(lines[i]); i += 1
            i += 1
            blocks.append({"t": "fence", "head": head, "body": body}); continue
        m = re.match(r"^(#{1,3})(!?) (.*)$", ln)
        if m:
            flush(); blocks.append({"t": "h", "level": len(m.group(1)), "plain": bool(m.group(2)), "text": m.group(3).strip()}); i += 1; continue
        if ln.strip() == "\\newpage":
            flush(); blocks.append({"t": "pagebreak"}); i += 1; continue
        if re.match(r"^\s*([-*•]|\d+\.)\s+", ln) and (not buf or True):
            flush()
            items = []
            while i < len(lines) and (re.match(r"^\s*([-•]|\d+\.)\s+", lines[i]) or (lines[i].startswith("   ") and items)):
                if re.match(r"^\s*([-•]|\d+\.)\s+", lines[i]):
                    items.append(lines[i].strip())
                else:
                    items[-1] += " " + lines[i].strip()
                i += 1
            blocks.append({"t": "list", "items": items}); continue
        if ln.strip() == "":
            flush(); i += 1; continue
        buf.append(ln); i += 1
    flush()
    return blocks

def parse_kv(head):
    """'fig key w=14 cols=2,1' -> (kind, key, {w:'14',...}) with quoted values allowed: cap="..." """
    toks = re.findall(r'(?:[^\s"]+="[^"]*")|\S+', head)
    kind, key = toks[0], toks[1] if len(toks) > 1 and "=" not in toks[1] else None
    kv = {}
    for t in toks[1:]:
        if "=" in t:
            k, v = t.split("=", 1); kv[k] = v.strip('"')
    return kind, key, kv

# ----------------------------------------------------------------------------- the document

class Builder:
    def __init__(self, pages=None):
        self.pages = pages or {}
        self.ctx = Ctx()
        self.doc = Document()
        self.chap = 0            # chapter counter
        self.chap_label = ""     # '3' or 'A'
        self.sec = 0; self.sub = 0
        self.counters = {"fig": 0, "tab": 0, "lst": 0, "alg": 0, "eq": 0}
        self.toc = []            # (level, label, title)
        self.lof = []; self.lot = []; self.lot_all = []
        self.setup_styles()

    # ---- document-level setup
    def setup_styles(self):
        d = self.doc
        st = d.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin, sec.right_margin = (Cm(2.5), Cm(2.2)) if SHORT else (Cm(3.0), Cm(2.5))
        sec.top_margin, sec.bottom_margin = (Cm(2.2), Cm(2.0)) if SHORT else (Cm(2.5), Cm(2.3))
        sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.1)
        cp = d.core_properties
        cp.title = "HF Model Finder: Constraint-Verified RAG for Hugging Face Models"
        cp.author = "Aditya Rane"
        cp.subject = "Capstone project report, MPSTME, SVKM's NMIMS"
        cp.keywords = "retrieval-augmented generation; Hugging Face; verification; Indic languages"

    def new_section(self, fmt=None, start=None, header=True, footer=True):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE) if self._first_done else self.doc.sections[0]
        self._first_done = True
        sectPr = s._sectPr
        for e in sectPr.findall(qn("w:pgNumType")): sectPr.remove(e)
        if fmt:
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), fmt)
            if start is not None: pg.set(qn("w:start"), str(start))
            sectPr.append(pg)
        for ref in sectPr.findall(qn("w:headerReference")) + sectPr.findall(qn("w:footerReference")):
            sectPr.remove(ref)      # a cloned section would otherwise share (and overwrite) the previous section's header/footer part
        s.header.is_linked_to_previous = False; s.footer.is_linked_to_previous = False
        for hp in list(s.header.paragraphs) + list(s.footer.paragraphs):
            for r in list(hp.runs): r._r.getparent().remove(r._r)
        if header:
            hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para_fmt(hp, WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=0)
            hp.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
            r = hp.add_run(SHORT_TITLE); set_font(r, 9, italic=True, color="404040")
            r = hp.add_run("\tMPSTME, SVKM's NMIMS"); set_font(r, 9, italic=True, color="404040")
            bottom_border(hp)
        if footer:
            fp = s.footer.paragraphs[0]
            para_fmt(fp, WD_ALIGN_PARAGRAPH.CENTER, line=1.0, after=0)
            add_field(fp, "PAGE", 10)
        return s

    # ---- front matter ----------------------------------------------------
    def centered(self, text, size=12, bold=False, italic=False, after=6, before=0, line=1.15, hl=False, keep_next=False):
        p = self.doc.add_paragraph()
        para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, line=line, before=before, after=after, keep_next=keep_next)
        add_inline(p, text, self.ctx, size=size, base_bold=bold, base_italic=italic)
        return p

    def logo(self, width=3.3, after=6):
        p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, line=1.0, after=after)
        p.add_run().add_picture(str(LOGO), width=Cm(width))

    def cover_simple(self):
        self.new_section(header=False, footer=False)
        c = self.centered
        c("HF Model Finder: Constraint Verified Retrieval-Augmented Generation for Finding Hugging Face Models", 20, True, after=18, before=10, line=1.2)
        c("Report submitted in partial fulfilment", 13, after=2)
        c("of", 13, after=2)
        c("B.Tech", 13, bold=True, after=2)
        c("in", 13, after=2)
        c("Computer Engineering", 13, bold=True, after=8)
        c("by", 13, after=4)
        c("**Harsh Bhandarkar** (B014), **Armaan Chopra** (B028),", 13, after=1)
        c("**Jay Gandhi** (B124), **Aditya Rane** (B129),", 13, after=1)
        c("**Yash Parab** (B179)", 13, after=12)
        c("Under the supervision of", 13, after=4)
        c("Dr. Archana Nanade", 14, bold=True, after=2)
        c("(Professor, Computer Engineering dept., MPSTME)", 12, after=22)
        c("SVKM's NMIMS University", 16, True, after=0)
        c("(Deemed-to-be University)", 12, after=8)
        self.logo(3.6, after=8)
        c("MUKESH PATEL SCHOOL OF TECHNOLOGY MANAGEMENT & ENGINEERING (MPSTME)", 13, True, after=2)
        c("Vile Parle (W), Mumbai-56", 12, True, after=16)
        c("(2026-27)", 13, True, after=0)

    def cover(self):
        if SIMPLE:
            return self.cover_simple()
        self.new_section(header=False, footer=False)
        c = self.centered
        c("HF Model Finder: Constraint-Verified Retrieval-Augmented Generation for Finding Hugging Face Models", 20, True, after=18, before=10, line=1.2)
        c("Capstone Project Report submitted in partial fulfilment", 13, after=2)
        c("of", 13, after=2)
        c("{{Name of the Programme}}", 13, bold=True, after=2)
        c("in", 13, after=2)
        c("{{Name of the Stream}}", 13, bold=True, after=8)
        c("by", 13, after=4)
        c("**Aditya Rane** ({{Roll No.}}, {{SAP ID}})", 14, after=12)
        c("Under the supervision of", 13, after=4)
        c("{{Name of Faculty Mentor}}", 14, bold=True, after=2)
        c("({{Designation, Name of the Department}}, MPSTME)", 12, after=22)
        c("SVKM's NMIMS University", 16, True, after=0)
        c("(Deemed-to-be University)", 12, after=8)
        self.logo(3.6, after=8)
        c("MUKESH PATEL SCHOOL OF TECHNOLOGY MANAGEMENT & ENGINEERING (MPSTME)", 13, True, after=2)
        c("Vile Parle (W), Mumbai-56", 12, True, after=16)
        c("({{Academic Year}})", 13, True, after=0)

    def certificate(self):
        self.new_section("lowerRoman", 1)
        c = self.centered
        p = c("CERTIFICATE", 16, True, after=10); p.runs[0].font.underline = True
        self.logo(3.0, after=12)
        p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5, after=10)
        add_inline(p, "This is to certify that the project entitled **\"HF Model Finder: Constraint-Verified Retrieval-Augmented Generation for Finding Hugging Face Models\"** has been done by **Mr. Aditya Rane** ({{Roll No.}}, {{SAP ID}}) under my guidance and supervision and has been submitted in partial fulfilment of the degree of {{Name of the Programme}} in {{Name of the Stream}} of MPSTME, SVKM's NMIMS (Deemed-to-be University), Mumbai, India.", self.ctx, 12)
        t = self.doc.add_table(rows=1, cols=2); no_borders(t); set_col_widths(t, [TEXT_W / 2, TEXT_W / 2])
        for cell, txt in zip(t.rows[0].cells, ["Project mentor (name and signature)<br>(Internal Guide)<br>{{Name of the mentor}}", "Examiner (name and signature)<br>{{Name of the examiner}}"]):
            cp = cell.paragraphs[0]; para_fmt(cp, WD_ALIGN_PARAGRAPH.LEFT, line=1.3, before=26, after=6); add_inline(cp, txt, self.ctx, 12)
        p = self.doc.add_paragraph(); para_fmt(p, line=1.5, before=24, after=0); add_inline(p, "**Date:** {{dd Month yyyy}}", self.ctx, 12)
        p = self.doc.add_paragraph(); para_fmt(p, line=1.5, after=0); add_inline(p, "**Place:** Mumbai", self.ctx, 12)
        p = self.doc.add_paragraph(); para_fmt(p, line=1.5, before=30, after=0); add_inline(p, "**(HoD)** (name and signature)<br>{{Name of the Head of Department}}", self.ctx, 12)

    def acknowledgement(self):
        self.page_break()
        p = self.centered("ACKNOWLEDGEMENT", 16, True, after=12)
        paras = [
            "I express my sincere gratitude to my mentor, {{Name of Faculty Mentor}}, for guidance, timely feedback and encouragement throughout this project, and to the Head of the Department, {{Name of the Head of Department}}, and the faculty of MPSTME, SVKM's NMIMS (Deemed-to-be University), for the academic environment that made this work possible.",
            "This work builds on the open resources of the Hugging Face Hub, whose model cards and metadata are the data of this study, and on free-tier services from Pinecone (vector search, hosted embeddings and re-ranking), Groq (language-model inference) and the open-source communities behind LangGraph, Streamlit and the Hugging Face libraries. I thank the authors of the papers cited in this report; the evaluation would not have been possible without their published methods.",
            "**Disclosure of tools.** The software was developed with the assistance of an AI coding assistant (Claude Code, Anthropic), which wrote and ran code under my direction. The relevance grades and the explanation-faithfulness scores reported in Chapter 4 were produced by independent language-model judges, not by human annotators; this is stated wherever those numbers appear. All design decisions, the choice of scope and the interpretation of results are my own responsibility.",
        ]
        for t in paras:
            p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5, after=8); add_inline(p, t, self.ctx, 12)
        p = self.doc.add_paragraph(); para_fmt(p, line=1.5, before=6, after=6); add_inline(p, "**Student:**", self.ctx, 12)
        t = self.doc.add_table(rows=2, cols=3); table_borders(t); cell_margins(t, 70, 70, 100, 100)
        set_col_widths(t, [7.5, 3.5, 4.5])
        hdr = ["NAME", "ROLL NO.", "SAP ID"]; row = ["Aditya Rane", "{{Roll No.}}", "{{SAP ID}}"]
        for i in range(3):
            c0 = t.rows[0].cells[i]; shade_cell(c0, "e7e6e6"); pp = c0.paragraphs[0]; para_fmt(pp, WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=0); add_inline(pp, hdr[i], self.ctx, 12, base_bold=True)
            c1 = t.rows[1].cells[i]; pp = c1.paragraphs[0]; para_fmt(pp, WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=0); add_inline(pp, row[i], self.ctx, 12)

    def abstract(self):
        self.page_break()
        self.centered("ABSTRACT", 16, True, after=12)
        for b in read_blocks(SRC / "abstract.md"):
            if b["t"] == "para":
                p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5, after=6); add_inline(p, b["text"], self.ctx, 12)

    def abstract_short(self):
        self.centered("ABSTRACT", 14, True, after=6)
        for b in read_blocks(SRC / "abstract.md"):
            if b["t"] == "para":
                p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.15, after=5); add_inline(p, b["text"], self.ctx, 11)
        self.centered("Note: the software was developed with the assistance of an AI coding assistant (Claude Code); see Section 3.8. Relevance and faithfulness grades were produced by AI judges, not humans.", 9, italic=True, after=4)

    def contents_short(self):
        self.centered("CONTENTS", 14, True, after=4, before=8)
        for (lvl, label, title) in self.all_headings:
            if lvl == 1:
                self.toc_line(lvl, label, title, self.pages.get("h:" + (label or title), ""))

    def page_break(self):
        p = self.doc.add_paragraph(); para_fmt(p, line=1.0, after=0)
        p.add_run().add_break(WD_BREAK.PAGE)

    # ---- generated lists (static, page numbers from the PDF pass)
    def toc_page(self):
        self.page_break()
        self.centered("TABLE OF CONTENTS", 16, True, after=10)
        p = self.doc.add_paragraph(); para_fmt(p, line=1.2, after=4)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run("Topics"); set_font(r, 12, True); r = p.add_run("\tPage"); set_font(r, 12, True)
        for key, label, title in [("lof", "", "List of Figures"), ("lot", "", "List of Tables"), ("abb", "", "Abbreviations")]:
            self.toc_line(0, label, title, self.pages.get("front:" + key, ""))
        for (lvl, label, title) in self.all_headings:
            self.toc_line(lvl, label, title, self.pages.get("h:" + (label or title), ""))

    def toc_line(self, lvl, label, title, page):
        p = self.doc.add_paragraph()
        indent = {0: 0.0, 1: 0.0, 2: 0.8, 3: 1.6}[lvl]
        para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, line=(1.0 if SHORT else 1.15), before=((3 if SHORT else 5) if lvl <= 1 else 0), after=(0 if SHORT else 1), left=indent)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        txt = (label + "  " if label else "") + title
        r = p.add_run(txt); set_font(r, 12 if lvl <= 1 else 11.5, bold=(lvl <= 1))
        r = p.add_run("\t" + str(page)); set_font(r, 12 if lvl <= 1 else 11.5, bold=(lvl <= 1))

    def list_page(self, key, title, col2, entries, prefix):
        self.page_break()
        self.centered(title, 16, True, after=10)
        t = self.doc.add_table(rows=1, cols=3); table_borders(t); cell_margins(t, 50, 50, 80, 80)
        widths = [2.3, 11.0, 2.2]
        for i, h in enumerate([prefix + " No.", col2, "Page No."]):
            c = t.rows[0].cells[i]; shade_cell(c, "e7e6e6"); pp = c.paragraphs[0]; para_fmt(pp, WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=0); add_inline(pp, h, self.ctx, 11, base_bold=True)
        row_props(t.rows[0], header=True)
        for label, cap in entries:
            cells = t.add_row().cells
            vals = [label.replace("Fig. ", "").replace("Table ", ""), cap, str(self.pages.get(f"{key}:{label}", ""))]
            for i, v in enumerate(vals):
                pp = cells[i].paragraphs[0]; para_fmt(pp, WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=0); add_inline(pp, v, self.ctx, 11)
            row_props(t.rows[-1])
        set_col_widths(t, widths)

    def abbreviations(self):
        self.page_break()
        self.centered("ABBREVIATIONS", 16, True, after=10)
        rows = [l for l in (SRC / "abbrev.md").read_text(encoding="utf-8").split("\n") if "|" in l]
        t = self.doc.add_table(rows=1, cols=2); table_borders(t); cell_margins(t, 40, 40, 80, 80)
        for i, h in enumerate(["Abbreviation", "Expansion"]):
            c = t.rows[0].cells[i]; shade_cell(c, "e7e6e6"); pp = c.paragraphs[0]; para_fmt(pp, WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=0); add_inline(pp, h, self.ctx, 11, base_bold=True)
        row_props(t.rows[0], header=True)
        for l in rows:
            a, b = [x.strip() for x in l.split("|", 1)]
            cells = t.add_row().cells
            for i, v in enumerate([a, b]):
                pp = cells[i].paragraphs[0]; para_fmt(pp, WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=0); add_inline(pp, v, self.ctx, 11, base_bold=(i == 0))
            row_props(t.rows[-1])
        set_col_widths(t, [3.6, 11.9])

    # ---- body -------------------------------------------------------------
    def heading(self, level, text, plain=False, appendix=False):
        d = self.doc
        if level == 1:
            self.counters = {k: 0 for k in self.counters}
            if plain:
                label, title = "", text
                full = text
                self.chap_label = ""
            elif appendix:
                letter = text.split(":")[0].split()[-1]
                self.chap_label = letter; label = ""; title = text; full = text
            else:
                self.chap += 1; self.chap_label = str(self.chap); self.sec = 0
                label = f"Chapter {self.chap}"; title = text; full = f"Chapter {self.chap}  {text}"
            self.sec = 0
            p = d.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT if (SHORT or label) else WD_ALIGN_PARAGRAPH.CENTER, line=1.15, before=(14 if SHORT else 0), after=(6 if SHORT else 14), keep_next=True)
            p.paragraph_format.page_break_before = not SHORT
            if SHORT:
                r = p.add_run(full); set_font(r, 15, True)
            elif label:
                r = p.add_run(label); set_font(r, 14, True); r.add_break(); r = p.add_run(title); set_font(r, 18, True)
            else:
                r = p.add_run(full); set_font(r, 18, True)
            bottom_border(p, "404040", 8)
            self.toc.append((1, label, title))
            self.all_headings_seen.append((1, label, title, label or title))
            return
        if level == 2:
            self.sec += 1; self.sub = 0
            lab = f"{self.chap_label}.{self.sec}" if self.chap_label else ""
            size, before = (12.5, 8) if SHORT else (14, 12)
        else:
            self.sub += 1
            lab = f"{self.chap_label}.{self.sec}.{self.sub}" if self.chap_label else ""
            size, before = (11, 6) if SHORT else (12, 8)
        p = d.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, line=1.15, before=before, after=(3 if SHORT else 6), keep_next=True)
        r = p.add_run((lab + "  " if lab else "") + text); set_font(r, size, True)
        if level == 2:
            self.toc.append((2, lab, text)); self.all_headings_seen.append((2, lab, text, lab))
        else:
            self.all_headings_seen.append((3, lab, text, lab))

    def paragraph(self, text):
        p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, line=LINE, after=(5 if SHORT else 6))
        add_inline(p, text, self.ctx, BODY)

    def bullets(self, items):
        n = 0
        for it in items:
            m = re.match(r"^(\d+)\.\s+(.*)$", it)
            if m:
                n += 1; mark, body = f"{m.group(1)}.", m.group(2)
            else:
                mark, body = "•", re.sub(r"^[-•]\s+", "", it)
            p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, line=LINE, after=3, hang=(1.0 if SHORT else 1.2, 0.6 if SHORT else 0.7))
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0 if SHORT else 1.2))
            r = p.add_run(mark + "\t"); set_font(r, BODY)
            add_inline(p, body, self.ctx, BODY)

    def next_label(self, kind):
        self.counters[kind] += 1
        return self.counters[kind]

    def figure(self, key, kv, body):
        n = self.next_label("fig")
        label = f"Fig. {self.chap_label}.{n}"
        w = float(kv.get("w", 14))
        items, cap = [], None
        for ln in body:
            if not ln.strip(): continue
            if "|" in ln and not ln.startswith("Caption") and Path(ln.split("|")[0].strip()).suffix in (".png", ".jpg", ".jpeg"):
                path, sub = [x.strip() for x in ln.split("|", 1)]
                items.append((path, sub))
            elif Path(ln.strip()).suffix in (".png", ".jpg", ".jpeg") and " " not in ln.strip():
                items.append((ln.strip(), ""))
            else:
                cap = ln.strip()
        if len(items) == 1:
            p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, line=1.0, before=6, after=4, keep_next=True)
            p.add_run().add_picture(str(self.resolve(items[0][0])), width=Cm(w))
        else:
            t = self.doc.add_table(rows=2, cols=len(items)); no_borders(t); cell_margins(t, 20, 20, 60, 60)
            each = TEXT_W / len(items)
            for i, (path, sub) in enumerate(items):
                c = t.rows[0].cells[i]; pp = c.paragraphs[0]; para_fmt(pp, WD_ALIGN_PARAGRAPH.CENTER, line=1.0, after=2, keep_next=True)
                pp.add_run().add_picture(str(self.resolve(path)), width=Cm(min(w, each - 0.3)))
                c2 = t.rows[1].cells[i]; pp = c2.paragraphs[0]; para_fmt(pp, WD_ALIGN_PARAGRAPH.CENTER, line=1.0, after=0, keep_next=True)
                add_inline(pp, sub, self.ctx, 10, base_italic=True)
            set_col_widths(t, [each] * len(items))
            row_props(t.rows[0]); row_props(t.rows[1])
        p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, line=1.0, before=2, after=12)
        r = p.add_run(label + ". "); set_font(r, 11, True)
        add_inline(p, cap or "", self.ctx, 11)
        self.lof.append((label, re.sub(r"\*|`", "", cap or "")))

    def resolve(self, path):
        for base in (HERE / "figs", HERE / "shots2" / "final", HERE, REPO):
            f = base / path
            if f.exists(): return f
            f = base / Path(path).name
            if f.exists() and base != REPO: return f
        raise FileNotFoundError(path)

    def table(self, key, kv, body):
        n = self.next_label("tab")
        label = f"Table {self.chap_label}.{n}"
        cap = body[0].strip()
        rows = [l for l in body[1:] if l.strip().startswith("|")]
        header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
        aligns = []
        for c in rows[1].strip().strip("|").split("|"):
            c = c.strip(); aligns.append("c" if c.startswith(":") and c.endswith(":") else "r" if c.endswith(":") else "l")
        data = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows[2:]]
        size = float(kv.get("size", 9 if SHORT else 10))
        ncol = len(header)
        if "cols" in kv:
            rel = [float(x) for x in kv["cols"].split(",")]
        else:
            rel = []
            for j in range(ncol):
                L = max([len(re.sub(r"[*`]", "", header[j]))] + [len(re.sub(r"[*`]", "", r[j])) for r in data])
                rel.append(min(max(L, 5), 42) ** 0.8)
        tot = sum(rel); width = float(kv.get("w", TEXT_W))
        widths = [width * r / tot for r in rel]
        p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, line=1.0, before=8, after=4, keep_next=True)
        r = p.add_run(label + ". "); set_font(r, 11, True)
        add_inline(p, cap, self.ctx, 11)
        self.lot_all.append((label, re.sub(r"\*|`", "", cap)))
        t = self.doc.add_table(rows=1, cols=ncol); table_borders(t); cell_margins(t, 35, 35, 70, 70)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        AL = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]; shade_cell(c, "e7e6e6"); pp = c.paragraphs[0]; para_fmt(pp, AL["l" if aligns[j] == "l" else "c"], line=1.0, after=0, keep_next=True)
            add_inline(pp, h, self.ctx, size, base_bold=True)
        row_props(t.rows[0], header=True)
        for ri, rdata in enumerate(data):
            cells = t.add_row().cells
            for j, v in enumerate(rdata):
                pp = cells[j].paragraphs[0]; para_fmt(pp, AL[aligns[j]], line=1.0, after=0, keep_next=(ri < len(data) - 1 and len(data) <= 12))
                add_inline(pp, v, self.ctx, size)
            row_props(t.rows[-1])
        set_col_widths(t, widths)
        p = self.doc.add_paragraph(); para_fmt(p, line=1.0, after=8)

    def code(self, key, kv, body, kind="lst"):
        n = self.next_label(kind)
        word = "Listing" if kind == "lst" else "Algorithm"
        label = f"{word} {self.chap_label}.{n}"
        cap = kv.get("cap", "")
        if "src" in kv:
            lines = (REPO / kv["src"]).read_text(encoding="utf-8").split("\n")
            if "lines" in kv:
                a, b = kv["lines"].split("-"); lines = lines[int(a) - 1:int(b)]
            code = lines
        else:
            code = body
        while code and not code[-1].strip(): code.pop()
        p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, line=1.0, before=8, after=3, keep_next=True)
        r = p.add_run(label + ". "); set_font(r, 11, True)
        add_inline(p, cap, self.ctx, 11)
        self.lot_all.append((label, re.sub(r"\*|`", "", cap)))
        size = float(kv.get("size", 7.5))
        for i, ln in enumerate(code):
            q = self.doc.add_paragraph(); para_fmt(q, WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=0, hang=(0.75, 0.55), widow=False)
            shade_par(q, "f3f3f3")
            q.paragraph_format.keep_together = True
            if len(code) <= 40 and i < len(code) - 1:
                q.paragraph_format.keep_with_next = True      # a short listing stays on one page
            txt = ln.replace("\t", "    ")
            r = q.add_run(txt if txt else " "); set_font(r, size, mono=True)
        q = self.doc.add_paragraph(); para_fmt(q, line=1.0, after=8)

    def equation(self, key, text):
        n = self.next_label("eq")
        label = f"({self.chap_label}.{n})"
        p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, line=1.3, before=4, after=6)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run("\t"); set_font(r, 12)
        add_inline(p, text, self.ctx, 12, base_italic=False)
        r = p.add_run("\t" + label); set_font(r, 12)

    def references(self):
        self.heading(1, "References", plain=True)
        for i, key in enumerate(self.ctx.cites, 1):
            p = self.doc.add_paragraph(); para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, line=(1.0 if SHORT else 1.15), after=(2 if SHORT else 5), hang=(1.0, 1.0))
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
            r = p.add_run(f"[{i}]\t"); set_font(r, 9 if SHORT else 11)
            add_inline(p, REFS[key], self.ctx, 9 if SHORT else 11)
        uncited = [k for k in REFS if k not in self.ctx.cites]
        if uncited:
            print("WARNING: references defined but not cited:", uncited)

    # ---- orchestration ---------------------------------------------------
    def render_block(self, b, *, appendix=False):
        t = b["t"]
        if t == "h":
            self.heading(b["level"], b["text"], plain=b["plain"], appendix=appendix and b["level"] == 1)
        elif t == "para":
            self.paragraph(b["text"])
        elif t == "list":
            self.bullets(b["items"])
        elif t == "pagebreak":
            self.page_break()
        elif t == "fence":
            kind, key, kv = parse_kv(b["head"])
            if kind == "fig": self.figure(key, kv, b["body"])
            elif kind == "table": self.table(key, kv, b["body"])
            elif kind == "code": self.code(key, kv, b["body"], "lst")
            elif kind == "algo": self.code(key, kv, b["body"], "alg")
            elif kind == "eq": self.equation(key, " ".join(x.strip() for x in b["body"]))
            else: raise ValueError(kind)

    def prepass(self, chapters):
        """Assign figure/table/listing numbers and citation numbers in document order (no output)."""
        ch, sec = 0, 0
        counters = {"fig": 0, "tab": 0, "lst": 0, "alg": 0, "eq": 0}
        label = ""
        for fname, appendix in chapters:
            for b in read_blocks(SRC / fname):
                if b["t"] == "h" and b["level"] == 1:
                    counters = {k: 0 for k in counters}
                    if b["plain"]: label = ""
                    elif appendix: label = b["text"].split(":")[0].split()[-1]
                    else: ch += 1; label = str(ch)
                if b["t"] == "para":
                    self.scan_cites(b["text"])
                elif b["t"] == "list":
                    for it in b["items"]: self.scan_cites(it)
                elif b["t"] == "fence":
                    kind, key, kv = parse_kv(b["head"])
                    word = {"fig": ("fig", "Fig."), "table": ("tab", "Table"), "code": ("lst", "Listing"), "algo": ("alg", "Algorithm"), "eq": ("eq", "Eq.")}[kind]
                    counters[word[0]] += 1
                    lab = f"{word[1]} {label}.{counters[word[0]]}".strip() if kind != "eq" else f"Eq. {label}.{counters['eq']}"
                    self.ctx.labels[f"{word[0]}:{key}"] = lab
                    for ln in b["body"]: self.scan_cites(ln)
                    self.scan_cites(kv.get("cap", ""))

    def scan_cites(self, text):
        for m in re.finditer(r"\[@([^\]]+)\]", text):
            for k in m.group(1).split(";"):
                self.ctx.cite_no(k.strip().lstrip("@"))

    def build(self, out_docx):
        chapters = [("ch1.md", False), ("ch2.md", False), ("ch3.md", False), ("ch4.md", False), ("ch5.md", False), ("ch6.md", False)]
        apps = [("appA.md", True)] if SIMPLE else [("appA.md", True), ("appB.md", True)] if SHORT else [("appA.md", True), ("appB.md", True), ("appC.md", True), ("appD.md", True)]
        self.prepass(chapters + [("refs_placeholder", False)] if False else chapters + apps)
        # headings list for the TOC comes from a dry run of body rendering order
        self._first_done = False
        self.all_headings_seen = []; self.lot_extra = []; self.lot_all = []
        # --- front matter
        if SHORT:
            self.cover(); self.new_section("lowerRoman", 1); self.abstract_short()
        else:
            self.cover(); self.certificate(); self.acknowledgement(); self.abstract()
        # We need headings and entries for the front-matter lists before rendering the body; collect them in a
        # throw-away builder run (same code path, so numbering matches exactly).
        dry = getattr(self, "_dry", None)
        if dry is None:
            d = Builder(self.pages); d._dry = True; d.dry_collect(chapters, apps)
            self.all_headings = [(l, lab, t) for (l, lab, t, _) in d.all_headings_seen if l <= 2]
            self.lof_entries, self.lot_entries = d.lof, d.lot_all
        if SHORT:
            self.contents_short()
        else:
            self.toc_page()
            self.list_page("lof", "LIST OF FIGURES", "Name of the figure", self.lof_entries, "Fig")
            self.list_page("lot", "LIST OF TABLES", "Name of the table / listing", self.lot_entries, "Table")
            self.abbreviations()
        # --- body
        self.new_section("decimal", 1)
        self.lof, self.lot, self.lot_extra, self.lot_all = [], [], [], []
        self.chap = 0; self.counters = {k: 0 for k in self.counters}
        for fname, _ in chapters:
            for b in read_blocks(SRC / fname): self.render_block(b)
        self.references()
        for fname, _ in apps:
            self.counters = {k: 0 for k in self.counters}
            for b in read_blocks(SRC / fname): self.render_block(b, appendix=True)
        fix_order(self.doc.element)
        for sec in self.doc.sections:
            for part in (sec.header, sec.footer): fix_order(part._element)
        self.doc.save(out_docx)

    def dry_collect(self, chapters, apps):
        """Render the body into a scratch document purely to learn headings, figure and table captions."""
        self.prepass(chapters + apps)
        self._first_done = True
        self.all_headings_seen = []; self.lot_extra = []; self.lot_all = []
        for fname, _ in chapters:
            for b in read_blocks(SRC / fname): self.render_block(b)
        self.heading(1, "References", plain=True)
        for fname, _ in apps:
            self.counters = {k: 0 for k in self.counters}
            for b in read_blocks(SRC / fname): self.render_block(b, appendix=True)


# ----------------------------------------------------------------------------- PDF pass and page lookup

def to_pdf(docx_path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(docx_path)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    return Path(docx_path).with_suffix(".pdf")

def page_texts(pdf):
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout).group(1))
    pages = []
    for i in range(1, n + 1):
        t = subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
        pages.append(t)
    return pages

def label_of(text):
    lines = [l.strip() for l in text.strip().split("\n") if l.strip()]
    return lines[-1] if lines else ""

def norm(s):
    return re.sub(r"\s+", " ", s).strip()

def collect_pages(b, pdf):
    texts = page_texts(pdf)
    out = {}
    body_start = None
    labels = [label_of(t) for t in texts]
    for i, lab in enumerate(labels):
        if lab == "1" and "Chapter 1" in texts[i] and "TABLE OF CONTENTS" not in texts[i]:
            body_start = i; break
    if body_start is None:
        print("could not locate body start"); return out, labels
    # front matter lists
    for i, t in enumerate(texts[:body_start]):
        n = norm(t)
        if "LIST OF FIGURES" in n and "front:lof" not in out: out["front:lof"] = labels[i]
        if "LIST OF TABLES" in n and "front:lot" not in out: out["front:lot"] = labels[i]
        if "ABBREVIATIONS Abbreviation Expansion" in n: out.setdefault("front:abb", labels[i])
    body = [norm(t) for t in texts[body_start:]]
    blabels = labels[body_start:]
    ptr = {"h": 0, "lof": 0, "lot": 0}
    def find(pat, kind):
        for i in range(ptr[kind], len(body)):
            if pat in body[i]:
                ptr[kind] = i; return blabels[i]
        return ""
    for (lvl, label, title, key) in b.all_headings_seen:
        if lvl == 3: continue
        if lvl == 1:
            pat = (label + " " + title) if label else title
            out["h:" + (label or title)] = find(norm(pat)[:60], "h")
        else:
            out["h:" + label] = find(norm(label + " " + title)[:55], "h")
    for label, cap in b.lof:
        out[f"lof:{label}"] = find(norm(label + ". " + cap)[:50], "lof")
    for label, cap in b.lot_all:
        out[f"lot:{label}"] = find(norm(label + ". " + cap)[:50], "lot")
    return out, labels


def main():
    OUT.mkdir(exist_ok=True)
    name = "HF_Model_Finder_Report_Simple" if SIMPLE else "HF_Model_Finder_Report_Short" if SHORT else "HF_Model_Finder_Capstone_Report"
    docx_path = OUT / f"{name}.docx"
    b = Builder({}); b.build(docx_path)
    pdf = to_pdf(docx_path)
    pages, labels = collect_pages(b, pdf)
    missing = [k for k, v in pages.items() if v == ""]
    print("pass 1: pages", len(labels), "missing", len(missing), missing[:10])
    for _ in range(2):
        b2 = Builder(pages); b2.build(docx_path)
        pdf = to_pdf(docx_path)
        new_pages, labels = collect_pages(b2, pdf)
        if new_pages == pages: break
        pages = new_pages
    missing = [k for k, v in pages.items() if v == ""]
    print("final: pages", len(labels), "missing", missing)
    json.dump(pages, open(OUT / "pages.json", "w"), indent=1)

if __name__ == "__main__":
    main()
